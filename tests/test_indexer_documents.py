"""
Tests de scripts/indexer_documents.py::indexer_document (la fonction reutilisable,
pas main() qui a besoin d'un vrai acces reseau hcp.ma). Chaine complete testee avec de
vraies bibliotheques (python-docx, sqlite3, chromadb, rank_bm25) et une fixture DOCX
generee a la volee (meme pattern que tests/test_extracteur.py et
tests/test_indexeur_texte.py) — seule la fonction d'embedding est factice (evite de
charger BGE-M3, ~2 Go).
"""
import re

import pytest

from src.base_donnees import connecter
from src.extracteur import Extracteur
from src.indexeur_texte import IndexeurTexte
from src.models import Document
from scripts.indexer_documents import indexer_document


def _fausse_fonction_embedding(textes: list[str]) -> list[list[float]]:
    vocabulaire = ["chomage", "emploi", "population"]
    vecteurs = []
    for texte in textes:
        mots = set(re.findall(r"\w+", texte.lower()))
        vecteur = [1.0 if mot in mots else 0.0 for mot in vocabulaire]
        if not any(vecteur):
            vecteur[0] = 0.01
        vecteurs.append(vecteur)
    return vecteurs


@pytest.fixture
def conn(tmp_path):
    connexion = connecter(tmp_path / "test.db")
    yield connexion
    connexion.close()


@pytest.fixture
def indexeur(tmp_path):
    return IndexeurTexte(dossier_chroma=tmp_path / "chroma", fonction_embedding=_fausse_fonction_embedding)


@pytest.fixture
def extracteur():
    return Extracteur()


def _document_docx(tmp_path, url="https://www.hcp.ma/test-ipc_a1111.html") -> Document:
    from docx import Document as DocxDocument

    docx = DocxDocument()
    docx.add_paragraph("Le taux de chomage a recule au premier trimestre 2026.")
    docx.add_paragraph("L'emploi progresse dans le secteur des services.")
    chemin = tmp_path / "note_ipc.docx"
    docx.save(chemin)

    return Document(
        id_document=None,
        url=url,
        titre="Note IPC test",
        date_publication="2026-07-20",
        langue="fr",
        categorie="Marche du travail",
        type="docx",
        texte_brut=str(chemin),
    )


def test_indexer_document_html_est_ignore_sans_toucher_la_base(conn, extracteur, indexeur):
    document_html = Document(
        id_document=None, url="https://www.hcp.ma/page.html", titre="Page",
        date_publication=None, langue="fr", categorie="Economie",
        type="html", texte_brut="<html><body><p>Contenu.</p></body></html>",
    )
    resume = indexer_document(conn, extracteur, indexeur, document_html)

    assert resume["statut"] == "ignore_html"
    assert resume["chunks"] == 0
    total_documents = conn.execute("SELECT COUNT(*) FROM document").fetchone()[0]
    assert total_documents == 0


def test_indexer_document_docx_insere_document_et_chunks(tmp_path, conn, extracteur, indexeur):
    document = _document_docx(tmp_path)
    resume = indexer_document(conn, extracteur, indexeur, document)

    assert resume["statut"] == "ok"
    assert resume["chunks"] >= 1
    assert resume["id_document"] is not None

    total_documents = conn.execute("SELECT COUNT(*) FROM document").fetchone()[0]
    assert total_documents == 1

    total_chunks = conn.execute("SELECT COUNT(*) FROM chunk").fetchone()[0]
    assert total_chunks == resume["chunks"]

    # Le chunk inséré en base doit être retrouvable par recherche hybride ensuite.
    resultats = indexeur.rechercher("chomage")
    assert len(resultats) >= 1


def test_indexer_document_docx_sans_texte_ne_cree_aucun_chunk(tmp_path, conn, extracteur, indexeur):
    from docx import Document as DocxDocument

    docx_vide = DocxDocument()  # aucun paragraphe
    chemin = tmp_path / "vide.docx"
    docx_vide.save(chemin)

    document = Document(
        id_document=None, url="https://www.hcp.ma/vide.html", titre="Vide",
        date_publication=None, langue="fr", categorie="Economie",
        type="docx", texte_brut=str(chemin),
    )
    resume = indexer_document(conn, extracteur, indexeur, document)

    assert resume["statut"] == "ok"
    assert resume["chunks"] == 0
    # Le document est quand meme insere (traçabilite), meme sans texte a indexer.
    total_documents = conn.execute("SELECT COUNT(*) FROM document").fetchone()[0]
    assert total_documents == 1


def test_indexer_document_meme_url_deux_fois_ne_duplique_pas_le_document(tmp_path, conn, extracteur, indexeur):
    document1 = _document_docx(tmp_path, url="https://www.hcp.ma/meme-url_a2222.html")
    document2 = _document_docx(tmp_path, url="https://www.hcp.ma/meme-url_a2222.html")

    indexer_document(conn, extracteur, indexeur, document1)
    indexer_document(conn, extracteur, indexeur, document2)

    total_documents = conn.execute("SELECT COUNT(*) FROM document").fetchone()[0]
    assert total_documents == 1


def test_indexer_document_meme_url_deux_fois_ne_duplique_pas_les_chunks(tmp_path, conn, extracteur, indexeur):
    """Regression du bug trouve le 27/08 (voir JOURNAL.md) : `document.url` est
    UNIQUE et empechait deja le doublon COTE DOCUMENT (voir test precedent), mais
    rien n'empechait de refaire extraction/chunking/embeddings et de reinserer des
    chunks en double a chaque relance de scripts/indexer_documents.py -- ce test
    couvrait deja l'URL, mais jamais le nombre de chunks, jusqu'a ce que ce bug soit
    trouve en conditions reelles (corpus jamais rafraichi car relancer le script
    l'aurait pollue)."""
    document1 = _document_docx(tmp_path, url="https://www.hcp.ma/meme-url_a3333.html")
    document2 = _document_docx(tmp_path, url="https://www.hcp.ma/meme-url_a3333.html")

    resume1 = indexer_document(conn, extracteur, indexeur, document1)
    resume2 = indexer_document(conn, extracteur, indexeur, document2)

    assert resume1["statut"] == "ok"
    assert resume2["statut"] == "deja_indexe"
    assert resume2["chunks"] == 0
    assert resume2["id_document"] == resume1["id_document"]

    total_chunks = conn.execute("SELECT COUNT(*) FROM chunk").fetchone()[0]
    assert total_chunks == resume1["chunks"]


def test_indexer_document_xlsx_avec_motif_reconnu_insere_les_indicateurs(tmp_path, conn, extracteur, indexeur):
    """Le repli PDF/XLSX (ConstructeurIndicateurs.structurer, voir son module) est
    appelé automatiquement ici sur tout document non-HTML -- ce test vérifie le
    branchement bout en bout (extraction -> structuration -> insertion SQLite) sur un
    XLSX reproduisant le motif réel (dictionnaire Code/Nom/Unité + table de données,
    voir data/raw/248688.xlsx)."""
    from openpyxl import Workbook

    classeur = Workbook()
    feuille_donnees = classeur.active
    feuille_donnees.title = "Données"
    feuille_donnees.append(["AN", "TRIM", "TX.CH.STR"])
    feuille_donnees.append([2025, 4, 11.5])

    feuille_meta = classeur.create_sheet("Métadonnées")
    feuille_meta.append(["Code", "Nom", "Unité"])
    feuille_meta.append(["AN", "Année", "-"])
    feuille_meta.append(["TRIM", "Trimestre", "-"])
    feuille_meta.append(["TX.CH.STR", "Taux de chômage strict", "%"])

    chemin = tmp_path / "indicateurs_test.xlsx"
    classeur.save(chemin)

    document = Document(
        id_document=None, url="https://www.hcp.ma/file/999999/", titre="Indicateurs test",
        date_publication="2025-11-01", langue="fr", categorie="Marche du travail",
        type="xlsx", texte_brut=str(chemin),
    )
    resume = indexer_document(conn, extracteur, indexeur, document)

    assert resume["statut"] == "ok"
    assert resume["indicateurs"] == 1

    ligne = conn.execute(
        "SELECT nom, valeur, unite, periode, region FROM indicateur WHERE id_document = ?",
        (resume["id_document"],),
    ).fetchone()
    assert ligne == ("Taux de chômage strict", 11.5, "%", "2025T4", None)
