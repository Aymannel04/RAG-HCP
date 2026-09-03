from unittest.mock import MagicMock, patch

from src.extracteur import Extracteur, _texte_lisible
from src.models import Document


def test_extraire_separe_texte_et_tableaux():
    html = """
    <html><body>
      <p>Premier paragraphe.</p>
      <p>Deuxieme paragraphe.</p>
      <table><tr><th>Annee</th><th>Valeur</th></tr><tr><td>2025</td><td>13.3</td></tr></table>
    </body></html>
    """
    doc = Document(
        id_document=None, url="https://example.test", titre="t",
        date_publication=None, langue="fr", categorie="Economie",
        type="html", texte_brut=html,
    )
    texte, tableaux = Extracteur().extraire(doc)

    assert "Premier paragraphe." in texte
    assert "Deuxieme paragraphe." in texte
    assert len(tableaux) == 1
    assert tableaux[0][0] == ["Annee", "Valeur"]
    assert tableaux[0][1] == ["2025", "13.3"]


# --- Garde-fou anti-texte-illisible (ajoute le 03/09, voir JOURNAL.md 30 aout et
# TODO.md) : PDF a police legacy/mise en page dense dont pdfplumber extrait des
# symboles bruts plutot que du texte reel -- cas deja trouve en base (document #75).

def test_texte_lisible_accepte_du_francais_normal():
    texte = (
        "Le taux de chômage a atteint 13,3 % au premier trimestre 2025, contre 12,9 % "
        "un trimestre plus tôt, selon les données du Haut-Commissariat au Plan."
    )
    assert _texte_lisible(texte) is True


def test_texte_lisible_rejette_des_symboles_bruts():
    # Simule le cas reel (document #75) : police legacy mal mappee, le texte extrait
    # est majoritairement des symboles/ponctuation sans rapport avec du texte reel.
    texte = "§¤#@%&*()_+=~^><{}[]|\\/§¤#@%&*()_+=~^><{}[]|\\/§¤#@%&*()_+=~^><{}[]|"
    assert _texte_lisible(texte) is False


def test_texte_lisible_conserve_par_defaut_si_trop_court():
    # Pas assez de signal sur un echantillon court -- mieux vaut garder que rejeter
    # a tort une page courte legitime (ex. page de garde avec peu de texte).
    assert _texte_lisible("§¤#@") is True
    assert _texte_lisible("") is True


def test_extraire_pdf_exclut_les_pages_illisibles_du_texte_mais_garde_leurs_tableaux():
    page_normale = MagicMock()
    page_normale.extract_text.return_value = (
        "Le taux de chômage a atteint 13,3 % selon les dernières données disponibles "
        "publiées par le Haut-Commissariat au Plan durant ce trimestre."
    )
    page_normale.extract_tables.return_value = []

    page_illisible = MagicMock()
    page_illisible.extract_text.return_value = (
        "§¤#@%&*()_+=~^><{}[]|\\/§¤#@%&*()_+=~^><{}[]|\\/§¤#@%&*()_+=~^><{}[]|"
    )
    page_illisible.extract_tables.return_value = [[["2025", "13.3"]]]

    pdf_factice = MagicMock()
    pdf_factice.pages = [page_normale, page_illisible]
    pdf_factice.__enter__.return_value = pdf_factice
    pdf_factice.__exit__.return_value = False

    with patch("pdfplumber.open", return_value=pdf_factice):
        texte, tableaux = Extracteur()._extraire_pdf("document_factice.pdf")

    assert "taux de chômage" in texte
    assert "§¤#" not in texte  # page illisible exclue du texte narratif
    assert tableaux == [[["2025", "13.3"]]]  # son tableau reste extrait normalement


def test_extraire_docx_separe_texte_et_tableaux(tmp_path):
    # ADR 0005 : support DOCX (pages IPC/IPPI, qui ne publient qu'en .docx). Fixture
    # generee a la volee avec python-docx, pas de fichier externe necessaire.
    from docx import Document as DocxDocument

    docx = DocxDocument()
    docx.add_paragraph("Premier paragraphe.")
    docx.add_paragraph("Deuxieme paragraphe.")
    docx.add_paragraph("")  # paragraphe vide : ne doit pas polluer le texte
    table = docx.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Annee"
    table.cell(0, 1).text = "Valeur"
    table.cell(1, 0).text = "2025"
    table.cell(1, 1).text = "13.3"

    chemin = tmp_path / "note_test.docx"
    docx.save(str(chemin))

    doc = Document(
        id_document=None, url="https://www.hcp.ma/attachment/999999/", titre="t",
        date_publication=None, langue="fr", categorie="Economie",
        type="docx", texte_brut=str(chemin),
    )
    texte, tableaux = Extracteur().extraire(doc)

    assert "Premier paragraphe." in texte
    assert "Deuxieme paragraphe." in texte
    assert texte.count("\n") == 1  # le paragraphe vide n'a pas laisse de ligne parasite
    assert len(tableaux) == 1
    assert tableaux[0][0] == ["Annee", "Valeur"]
    assert tableaux[0][1] == ["2025", "13.3"]
